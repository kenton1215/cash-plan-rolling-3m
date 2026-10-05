"""Excel 导出: 三个月滚动资金计划工作簿 (openpyxl)

工作表: 1.资金计划汇总(经典主表) 2.现金流入明细 3.现金流出明细 4.缺口诊断与建议
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
SECTION_FONT = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
SECTION_FILL = PatternFill("solid", fgColor="E7EEF8")
NORMAL_FONT = Font(name="微软雅黑", size=10)
BOLD_FONT = Font(name="微软雅黑", size=10, bold=True)
RED_FONT = Font(name="微软雅黑", size=10, bold=True, color="C00000")
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
RED_FILL = PatternFill("solid", fgColor="FDE9E9")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _cat_order(plan, kind):
    cats = []
    for r in plan["rows"]:
        for d in r["details"].get(kind, []):
            if d.get("category") not in cats:
                cats.append(d["category"])
    return cats


def _cat_values(plan, kind, cat):
    return [round(sum(d.get("amount", 0) for d in r["details"].get(kind, [])
                      if d.get("category") == cat), 2) for r in plan["rows"]]


def _style_header(ws, row, ncol):
    for j in range(1, ncol + 1):
        c = ws.cell(row, j)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = CENTER
    ws.row_dimensions[row].height = 30


def _fill_detail_sheet(ws, plan, kind):
    headers = ["期间", "类别", "项目", "金额(万元)", "说明"]
    for j, h in enumerate(headers, 1):
        ws.cell(1, j, h)
    _style_header(ws, 1, len(headers))
    ri = 2
    for r in plan["rows"]:
        for d in r["details"].get(kind, []):
            vals = [r["month"], d.get("category", ""), d.get("item", ""),
                    round(d.get("amount", 0), 2), d.get("note", "")]
            for j, v in enumerate(vals, 1):
                c = ws.cell(ri, j, v)
                c.font = NORMAL_FONT
                c.border = BORDER
                c.alignment = LEFT if j in (3, 5) else CENTER
            ri += 1
    widths = [12, 16, 34, 14, 40]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def build_excel(company_name, plan, diagnosis_text="", issues=None):
    """生成完整工作簿, 返回 xlsx 字节"""
    rows = plan["rows"]
    months = [f"{r['month']}\n({'实际' if r['is_actual'] else '预测'})" for r in rows]
    ncol = len(rows) + 1
    wb = Workbook()

    # ================= 1. 资金计划汇总 =================
    ws = wb.active
    ws.title = "资金计划汇总"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    t = ws.cell(1, 1, f"{company_name} · 三个月滚动资金计划（编制月 {plan['plan_month']}）")
    t.font = TITLE_FONT
    t.alignment = CENTER
    ws.row_dimensions[1].height = 32

    ws.cell(2, 1, "项目（单位：万元）")
    for j, m in enumerate(months):
        ws.cell(2, j + 2, m)
    _style_header(ws, 2, ncol)

    line_defs = []
    line_defs.append(("期初现金余额", [r["opening"] for r in rows], "normal"))
    line_defs.append(("一、现金流入", [""] * len(rows), "section"))
    for cat in _cat_order(plan, "inflows"):
        line_defs.append((f"　{cat}", _cat_values(plan, "inflows", cat), "normal"))
    line_defs.append(("　现金流入合计", [r["inflow"] for r in rows], "total"))
    line_defs.append(("二、现金流出", [""] * len(rows), "section"))
    for cat in _cat_order(plan, "outflows"):
        line_defs.append((f"　{cat}", _cat_values(plan, "outflows", cat), "normal"))
    line_defs.append(("　现金流出合计", [r["outflow"] for r in rows], "total"))
    line_defs.append(("三、现金净流量", [r["net"] for r in rows], "total"))
    line_defs.append(("四、期末现金余额", [r["closing"] for r in rows], "closing"))
    line_defs.append(("五、现金安全底线", [r["safety_line"] for r in rows], "normal"))
    line_defs.append(("六、资金缺口（低于底线金额）", [r["gap"] for r in rows], "gap"))

    ri = 3
    for name, vals, kind in line_defs:
        c0 = ws.cell(ri, 1, name)
        c0.border = BORDER
        c0.alignment = LEFT
        if kind == "section":
            c0.font = SECTION_FONT
            c0.fill = SECTION_FILL
        elif kind == "total":
            c0.font = BOLD_FONT
        else:
            c0.font = NORMAL_FONT
        for j, v in enumerate(vals):
            c = ws.cell(ri, j + 2, "" if v == "" else round(float(v), 2))
            c.border = BORDER
            c.alignment = CENTER
            if kind == "total":
                c.font = BOLD_FONT
                c.fill = TOTAL_FILL
            elif kind == "gap" and isinstance(v, (int, float)) and v > 0:
                c.font = RED_FONT
                c.fill = RED_FILL
            elif kind == "closing" and isinstance(v, (int, float)) and v < plan["safety_line"]:
                c.font = RED_FONT
                c.fill = RED_FILL
            elif kind == "section":
                c.fill = SECTION_FILL
            else:
                c.font = NORMAL_FONT
        ri += 1

    ws.column_dimensions["A"].width = 30
    for j in range(2, ncol + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16

    # ================= 2/3. 流入/流出明细 =================
    ws2 = wb.create_sheet("现金流入明细")
    _fill_detail_sheet(ws2, plan, "inflows")
    ws3 = wb.create_sheet("现金流出明细")
    _fill_detail_sheet(ws3, plan, "outflows")

    # ================= 4. 缺口诊断与建议 =================
    ws4 = wb.create_sheet("缺口诊断与建议")
    ws4.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    c = ws4.cell(1, 1, "智能诊断与应对建议（DeepSeek 大模型 + 规则引擎）")
    c.font = TITLE_FONT
    ri = 3
    text = (diagnosis_text or "未生成诊断, 请在系统中点击「生成智能诊断」。")
    for line in text.splitlines():
        if not line.strip():
            continue
        c = ws4.cell(ri, 1, line)
        c.font = NORMAL_FONT
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws4.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=2)
        ws4.row_dimensions[ri].height = max(18, 16 * (len(line) // 60 + 1))
        ri += 1
    if issues:
        ri += 1
        c = ws4.cell(ri, 1, "数据校验问题:")
        c.font = BOLD_FONT
        ri += 1
        for it in issues[:20]:
            c = ws4.cell(ri, 1, f"- [{it['level']}] {it['module']}: {it['message']}")
            c.font = NORMAL_FONT
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            ws4.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=2)
            ws4.row_dimensions[ri].height = max(18, 16 * (len(it["message"]) // 60 + 1))
            ri += 1
    ws4.column_dimensions["A"].width = 50
    ws4.column_dimensions["B"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
