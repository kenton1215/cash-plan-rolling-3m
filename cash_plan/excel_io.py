"""Excel 模板下载与一键导入(按业务模块)

流程: 按业务模块分别下载模板 → 发给对应部门离线填写 → 资金计划专员按模块上传一键导入
导入规则:
  - 按所选模块和角色权限过滤工作表(其他模块/无权限的表跳过并提示)
  - 逐行校验: 金额/比例0~1/月份YYYY-MM/日期YYYY-MM-DD/是-否/borrow-repay
  - 错误行不导入并给出"第几行-哪个字段-原因"
  - 普通表: 追加模式, 与库中完全相同的行自动跳过(重复导入安全)
  - 工资计划/期初现金/月度实际数: 按(企业,月份)更新覆盖, 月份为必填
"""
import io
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .database import session_scope
from .schema import MODULES_BY_KEY, TABLE_BY_SHEET, TABLES, UNIQUE_MONTH_TABLES

MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
NOTE_FONT = Font(name="微软雅黑", size=11)
BOLD_NOTE = Font(name="微软雅黑", size=11, bold=True, color="1F4E79")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

_WIDTHS = {"text": 22, "num": 16, "int": 10, "bool": 12, "rate": 16, "month": 16, "date": 16}


def allowed_tables(role, module_key=None):
    """该角色有权维护的表; module_key 给定时只返回该模块的表"""
    return [t for t in TABLES
            if role in MODULES_BY_KEY[t[0]][2] and (module_key is None or t[0] == module_key)]


def module_name(module_key):
    """模块标签去掉emoji后的中文名(用于文件名)"""
    label = MODULES_BY_KEY[module_key][1]
    return label.split(" ", 1)[1] if " " in label else label


# ---------------------------------------------------------------------------
# 模板下载
# ---------------------------------------------------------------------------
def build_template(role, module_key=None):
    """生成Excel填报模板(按业务模块): 「填写说明」+ 该模块每张表一个工作表"""
    wb = Workbook()
    intro = wb.active
    intro.title = "填写说明"
    lines = [
        "【资金计划数据填报模板 · 填写说明】",
        "",
        "一、通用规则",
        "1. 所有金额单位为万元; 日期格式 YYYY-MM-DD(如2026-10-15), 月份格式 YYYY-MM(如2026-10)。",
        "2. 比例类字段(回款率/赢单概率/兑付概率/预付比例)填0~1之间的小数, 如95%填0.95。",
        "3. 勾选类字段(是否付现/研发试产/预测订单)填: 是 或 否。",
        "4. 贷款与筹资表中「类型」填 borrow(新增借款) 或 repay(还款)。",
        "5. 只填写需要的行, 空白行自动忽略; 不要修改、增删表头文字, 不要改工作表名称。",
        "6. 填写完成后交给资金计划专员「一键导入」; 导入为追加模式, 重复行自动跳过, 可放心重复导入。",
        "",
        "二、本模板包含的工作表",
    ]
    for _, sheet, _model, fields, hint in allowed_tables(role, module_key):
        cols = "、".join(f[1] for f in fields)
        lines.append(f"· 【{sheet}】{hint}")
        lines.append(f"  列: {cols}")
    for i, line in enumerate(lines, 1):
        c = intro.cell(i, 1, line)
        c.font = BOLD_NOTE if (i == 1 or line.startswith("·")) else NOTE_FONT
    intro.column_dimensions["A"].width = 110

    for _, sheet, _model, fields, _hint in allowed_tables(role, module_key):
        ws = wb.create_sheet(sheet)
        for j, (_attr, label, kind) in enumerate(fields, 1):
            c = ws.cell(1, j, label)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = _WIDTHS.get(kind, 16)
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 一键导入
# ---------------------------------------------------------------------------
def import_excel(file_bytes, company_id, role, module_key=None):
    """按业务模块解析并导入Excel, 返回逐表报告 [{sheet, added, updated, skipped, errors}]"""
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception:
        return [{"sheet": "文件", "added": 0, "updated": 0, "skipped": 0,
                 "errors": ["文件无法解析, 请确认上传的是系统导出的xlsx模板"]}]
    allowed = {t[1] for t in allowed_tables(role, module_key)}
    report = []
    with session_scope() as sess:
        for sheet_name in wb.sheetnames:
            if sheet_name == "填写说明":
                continue
            t = TABLE_BY_SHEET.get(sheet_name)
            if t is None:
                report.append({"sheet": sheet_name, "added": 0, "updated": 0, "skipped": 0,
                               "errors": ["未识别的数据表, 已跳过(请使用系统导出的模板)"]})
                continue
            if module_key is not None and t[0] != module_key:
                report.append({"sheet": sheet_name, "added": 0, "updated": 0, "skipped": 0,
                               "errors": [f"该工作表不属于所选模块「{MODULES_BY_KEY[module_key][1]}」, "
                                          f"已跳过(请确认模块选择是否正确)"]})
                continue
            if sheet_name not in allowed:
                report.append({"sheet": sheet_name, "added": 0, "updated": 0, "skipped": 0,
                               "errors": ["当前角色无权限导入该表, 已跳过"]})
                continue
            report.append(_import_sheet(sess, wb[sheet_name], t, company_id))
    return report


def _import_sheet(sess, ws, t, company_id):
    _module, sheet, model, fields, _hint = t
    # 表头校验
    header = [ws.cell(1, j).value for j in range(1, len(fields) + 1)]
    if [(str(h).strip() if h is not None else "") for h in header] != [f[1] for f in fields]:
        return {"sheet": sheet, "added": 0, "updated": 0, "skipped": 0,
                "errors": ["表头与模板不一致, 请使用系统导出的模板"]}
    # 逐行解析
    errors, data_rows = [], []
    for ri, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all((v is None or str(v).strip() == "") for v in row):
            continue
        data, errs = _parse_row(fields, row, ri)
        if errs:
            errors.extend(errs)
            continue
        if model in UNIQUE_MONTH_TABLES and not data.get("month"):
            errors.append(f"第{ri}行「月份」: 月份为必填(YYYY-MM), 该行未导入")
            continue
        if all((data[f[0]] in (None, "", 0, 0.0, False)) for f in fields):
            continue
        data_rows.append(data)
    if not data_rows and not errors:
        return {"sheet": sheet, "added": 0, "updated": 0, "skipped": 0,
                "errors": [], "empty": True}

    added = updated = skipped = 0
    if model in UNIQUE_MONTH_TABLES:
        # 按(企业,月份)更新覆盖, 空单元格保留原值
        for data in data_rows:
            obj = sess.query(model).filter(model.company_id == company_id,
                                           model.month == data["month"]).first()
            if obj:
                for k, v in data.items():
                    if v not in (None, "", False):
                        setattr(obj, k, v)
                updated += 1
            else:
                sess.add(model(company_id=company_id, **data))
                added += 1
    else:
        # 追加模式: 与库中(含本次已导入)完全相同的行自动跳过
        existing = sess.query(model).filter(model.company_id == company_id).all()
        for data in data_rows:
            if any(all(data[k] == getattr(o, k) for k in data) for o in existing):
                skipped += 1
                continue
            obj = model(company_id=company_id, **data)
            sess.add(obj)
            existing.append(obj)
            added += 1
    return {"sheet": sheet, "added": added, "updated": updated, "skipped": skipped,
            "errors": errors[:20]}


def _parse_row(fields, row, ri):
    data, errs = {}, []
    for (attr, label, kind), v in zip(fields, row):
        try:
            data[attr] = _convert(kind, v)
        except ValueError as e:
            errs.append(f"第{ri}行「{label}」: {e}")
    return data, errs


def _convert(kind, v):
    if v is None or (isinstance(v, str) and not v.strip()):
        if kind in ("num", "int", "rate"):
            return None
        return False if kind == "bool" else ""
    if kind in ("num", "int", "rate"):
        if isinstance(v, (int, float)):
            fv = float(v)
        else:
            try:
                fv = float(str(v).strip().replace(",", ""))
            except ValueError:
                raise ValueError(f"必须是数字, 实际为「{v}」")
        if kind == "int":
            if fv != int(fv):
                raise ValueError(f"必须是整数, 实际为「{v}」")
            return int(fv)
        if kind == "rate":
            if not 0 <= fv <= 1:
                raise ValueError(f"必须在0~1之间, 实际为「{v}」")
        return fv
    if kind == "bool":
        if isinstance(v, bool):
            return v
        s = str(v).strip().lower()
        if s in ("是", "true", "1", "y", "yes", "√"):
            return True
        if s in ("否", "false", "0", "n", "no", "×", "x"):
            return False
        raise ValueError(f"请填「是/否」, 实际为「{v}」")
    if kind == "month":
        if isinstance(v, datetime):
            return v.strftime("%Y-%m")
        s = str(v).strip()
        if not MONTH_RE.match(s):
            raise ValueError(f"月份格式应为YYYY-MM, 实际为「{v}」")
        return s
    if kind == "date":
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d")
        s = str(v).strip()
        if not DATE_RE.match(s):
            raise ValueError(f"日期格式应为YYYY-MM-DD, 实际为「{v}」")
        return s
    return str(v).strip()
